import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
import os

class InfotiposProcessor:

    def __init__(self, input_file: str, expected_columns: str, infotipo: str, case: str, output_directory: str):

        self.input_file = input_file
        self.expected_columns = expected_columns
        self.infotipo = infotipo
        self.case = case
        self.output_directory = output_directory


    def process_infotipos(self)->bool:
        try:
            # Validate input file exists
            if not os.path.exists(self.input_file):
                print(f"Error: Input file {self.input_file} does not exist.")
                return False

            # Process column configuration
            columns = [col.strip("'").lower() for col in self.expected_columns.split(",")]

            # Remove 'cvacio' placeholders to check actual columns
            actual_columns = [col for col in columns if col != 'cvacio']

            # Read Excel file
            try:
                df = pd.read_excel(self.input_file)
            except Exception as e:
                print(f"Error reading Excel file: {e}")
                return False

            # Normalize column names
            df.columns = [str(col).lower().strip() for col in df.columns]
            df = df.astype(str)

            # Validate required columns
            missing_columns = set(actual_columns) - set(df.columns)
            if missing_columns:
                print(f"Missing columns: {missing_columns }")
                return False

            # Select and prepare DataFrame
            df_processed = df[actual_columns]

            # Handle 'cvacio' placeholders
            cvacio_positions = [i for i, col in enumerate(columns) if col == 'cvacio']
            for i, pos in enumerate(cvacio_positions):
                df_processed.insert(pos, f"vacio{i+1}", "")

            # Create output file path
            output_file = f"{self.infotipo}_{self.case}.xlsx"
            full_output_path = os.path.join(self.output_directory, output_file)

            # Create new workbook and write data
            wb = Workbook()
            ws = wb.active

            # Write data starting from row 7 (fila_inicio = 6)
            rows = dataframe_to_rows(df_processed, index=False, header=True)
            for i, row in enumerate(rows, start=7):
                for j, value in enumerate(row, start=1):
                    ws.cell(row=i, column=j, value=value)

            # Save the workbook
            wb.save(full_output_path)
            print(f"Successfully processed file. Output saved to: {full_output_path}")
            return True

        except Exception as error:
            print(f"Unexpected error during processing: {error}")
            return False
